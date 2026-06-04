"""Unit tests for text extraction (TXT, MD, DOCX, PDF, unsupported)."""
from pathlib import Path

import pytest
from docx import Document as Docx

from app.services.text_extraction import extract_text


def test_txt_round_trip(tmp_path: Path):
    p = tmp_path / "hello.txt"
    p.write_text("Hello, world.\nSecond line.", encoding="utf-8")
    assert "Hello, world" in extract_text(p)


def test_md_treated_as_text(tmp_path: Path):
    p = tmp_path / "doc.md"
    p.write_text("# Heading\n\nSome paragraph.", encoding="utf-8")
    text = extract_text(p)
    assert "Heading" in text and "paragraph" in text


def test_unicode_preserved(tmp_path: Path):
    p = tmp_path / "uk.txt"
    p.write_text("Привіт, світе!", encoding="utf-8")
    assert "Привіт" in extract_text(p)


def test_unsupported_extension_raises(tmp_path: Path):
    p = tmp_path / "binary.exe"
    p.write_bytes(b"MZ\x90\x00")
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(p)


def test_docx_paragraphs_and_table(tmp_path: Path):
    p = tmp_path / "sample.docx"
    d = Docx()
    d.add_paragraph("First paragraph.")
    d.add_paragraph("Second paragraph.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Name"
    t.rows[0].cells[1].text = "Value"
    t.rows[1].cells[0].text = "Alpha"
    t.rows[1].cells[1].text = "42"
    d.save(p)

    text = extract_text(p)
    assert "First paragraph" in text
    assert "Second paragraph" in text
    # Table cells are joined with " | "
    assert "Name | Value" in text
    assert "Alpha | 42" in text


def test_docx_skips_empty_paragraphs(tmp_path: Path):
    p = tmp_path / "sparse.docx"
    d = Docx()
    d.add_paragraph("")
    d.add_paragraph("Actual content.")
    d.add_paragraph("   ")
    d.save(p)
    text = extract_text(p)
    assert text.strip() == "Actual content."


def test_corrupted_pdf_raises(tmp_path: Path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4 not a real pdf body \x00\x01\x02")
    # pypdf raises on a malformed file; we just want it to surface (caller marks doc failed)
    with pytest.raises(Exception):
        extract_text(p)


def test_extension_case_insensitive(tmp_path: Path):
    p = tmp_path / "doc.TXT"
    p.write_text("upper-ext content", encoding="utf-8")
    assert "upper-ext" in extract_text(p)
