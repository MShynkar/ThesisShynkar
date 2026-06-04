"""Extract plain text from PDF, DOCX, TXT, and Markdown."""
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


def extract_text(file_path: str | Path, mime_type: str | None = None) -> str:
    """Dispatch to the right extractor based on extension."""
    path = Path(file_path)
    suffix = path.suffix.lower().lstrip(".")

    if suffix == "pdf":
        return _extract_pdf(path)
    if suffix == "docx":
        return _extract_docx(path)
    if suffix in ("txt", "md", "markdown"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # Some PDFs have malformed pages; skip them rather than failing the whole upload
            continue
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
